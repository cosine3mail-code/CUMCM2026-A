# -*- coding: utf-8 -*-
"""问题1预热平衡阶段求解器。

模型与物性来源：附录2；关键参数：h=25、h_m=8e-7、Δt=0.25 s、Δr=2.5×10⁻⁵ m。
输入数据：``../附件/附件1.xlsx``；输出位置：本目录 ``结果/``。
数值方法：有限体积 + 后向 Euler + Thomas 三对角 + Picard 迭代。
运行方式：在项目根目录执行 ``python 问题1/问题1_求解器.py``。
"""
import os
import numpy as np
from scipy.interpolate import PchipInterpolator
import openpyxl

# ===================== 路径 =====================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
RESULT_DIR = os.path.join(SCRIPT_DIR, '结果')
ATT1 = os.path.join(ROOT, '附件', '附件1.xlsx')

# ===================== 参数（附录2） =====================
R    = 0.02        # 半径 m
rho  = 820.0       # 密度 kg/m^3
cp   = 2600.0      # 比热容 J/(kg·K)
k    = 0.36        # 热传导系数 W/(m·K)
h    = 25.0        # 对流换热系数 W/(m^2·K)
hm   = 8e-7        # 对流传质系数 m/s
D0   = 7e-9        # 扩散系数前置因子 m^2/s
Da   = 0.89        # D = D0 * exp(-Da / C)

T0   = 28.0        # 初始温度 ℃
C0   = 2.55        # 初始水分浓度 kg/kg

t_end = 1800.0     # 求解总时长 s
dt    = 0.25       # 时间步长 s
N     = 800        # 径向网格数
dr    = R / N      # 与论文算法1一致（N_r=800、Δt=0.25 s、Δr=2.5×10⁻⁵ m）
alpha = k / (rho * cp)   # 热扩散率 m^2/s


def Dfun(C):
    """水分浓度扩散系数（附录2）：D = D0 * exp(-Da / C)"""
    return D0 * np.exp(-Da / C)


# ===================== 读取附件1 =====================
def read_attachment1(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]   # 跳过表头
    t = np.array([r[0] for r in rows], dtype=float)
    T = np.array([r[1] for r in rows], dtype=float)
    C = np.array([r[2] for r in rows], dtype=float)
    return t, T, C

t_air, T_air, C_air = read_attachment1(ATT1)
T_air_interp = PchipInterpolator(t_air, T_air)
C_air_interp = PchipInterpolator(t_air, C_air)

# ===================== 网格与输出节点 =====================
r = np.arange(N + 1) * dr
n_steps = int(round(t_end / dt))

idx_table = [0, 200, 400, 600, 800]        # 0、0.5、1、1.5、2 cm
idx_full  = np.arange(0, N + 1, int(0.001 / dr))  # 每 0.1 cm，共 21 点
r_full    = idx_full * dr                 # 对应距离 m
r_table   = np.array(idx_table) * dr


# ===================== Thomas 算法 =====================
def thomas(a, b, c, d):
    """解三对角方程组 a[i]x[i-1]+b[i]x[i]+c[i]x[i+1]=d[i]，a[0]=c[N]=0"""
    a = a.copy(); b = b.copy(); c = c.copy(); d = d.copy()
    n = len(b)
    cp_ = np.zeros(n); dp_ = np.zeros(n)
    cp_[0] = c[0] / b[0]
    dp_[0] = d[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp_[i - 1]
        cp_[i] = c[i] / m
        dp_[i] = (d[i] - a[i] * dp_[i - 1]) / m
    x = np.zeros(n)
    x[-1] = dp_[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp_[i] - cp_[i] * x[i + 1]
    return x


# ===================== 热方程系数（线性，常数矩阵） =====================
s   = alpha * dt / dr**2
A_h = 2 * N * h * dt / ((N - 0.25) * rho * cp * dr)   # 表面对流项系数
B_h = 2 * (N - 0.5) * s / (N - 0.25)

aT = np.zeros(N + 1); bT = np.zeros(N + 1); cT = np.zeros(N + 1)
bT[0] = 1 + 4 * s;  cT[0] = -4 * s                       # 中心
for i in range(1, N):
    aT[i] = -(i - 0.5) / i * s
    cT[i] = -(i + 0.5) / i * s
    bT[i] = 1 + 2 * s
aT[N] = -B_h;  bT[N] = 1 + A_h + B_h;  cT[N] = 0          # 表面(半格点)

# ===================== 水分方程表面对流项系数（常数） =====================
A_m = 2 * N * hm * dt / ((N - 0.25) * dr)


# ===================== 时间推进 =====================
T = np.full(N + 1, T0)
C = np.full(N + 1, C0)

T_full = np.zeros((n_steps, len(idx_full)))   # 每隔 0.25 s、每隔 0.1 cm 记录
C_full = np.zeros((n_steps, len(idx_full)))

times_table = [100, 300, 600, 900, 1200, 1500, 1800]
table_map = {t: i for i, t in enumerate(times_table)}
T_table = np.zeros((len(times_table), len(idx_table)))
C_table = np.zeros((len(times_table), len(idx_table)))

for n in range(n_steps):
    t_new = (n + 1) * dt

    # ---- 热方程（线性，直接求解一次） ----
    dT = T.copy()
    dT[N] += A_h * T_air_interp(t_new)
    T = thomas(aT, bT, cT, dT)

    # ---- 水分方程（Picard 迭代） ----
    C_new = C.copy()
    for it in range(50):
        C_prev = C_new.copy()
        C_face = 0.5 * (C_prev[:-1] + C_prev[1:])        # 界面水分浓度
        lam = Dfun(C_face) * dt / dr**2                  # 界面扩散数 λ_i

        aC = np.zeros(N + 1); bC = np.zeros(N + 1); cC = np.zeros(N + 1); dC = np.zeros(N + 1)
        bC[0] = 1 + 4 * lam[0];  cC[0] = -4 * lam[0];  dC[0] = C[0]     # 中心
        for i in range(1, N):
            aC[i] = -(i - 0.5) / i * lam[i - 1]
            cC[i] = -(i + 0.5) / i * lam[i]
            bC[i] = 1 + (i - 0.5) / i * lam[i - 1] + (i + 0.5) / i * lam[i]
            dC[i] = C[i]
        B_m = 2 * (N - 0.5) * lam[N - 1] / (N - 0.25)                    # 表面
        aC[N] = -B_m;  bC[N] = 1 + A_m + B_m;  cC[N] = 0
        dC[N] = C[N] + A_m * C_air_interp(t_new)

        C_new = thomas(aC, bC, cC, dC)
        if np.max(np.abs(C_new - C_prev)) < 1e-12:
            break
    C = C_new

    # ---- 记录 ----
    T_full[n] = T[idx_full]
    C_full[n] = C[idx_full]
    if t_new in table_map:
        j = table_map[t_new]
        T_table[j] = T[idx_table]
        C_table[j] = C[idx_table]

# ===================== 输出表1/表2（论文表格） =====================
def print_table(title, times, r_cm, data):
    print(f"\n{title}")
    header = "时间/s | " + " | ".join(f"{v:g}cm" for v in r_cm)
    print(header)
    print("-" * len(header))
    for i, t in enumerate(times):
        line = f"{t:>6d} | " + " | ".join(f"{data[i, j]:.4f}" for j in range(len(r_cm)))
        print(line)

r_table_cm = r_table * 100
print_table("表1  30分钟内药材的温度（℃）", times_table, r_table_cm, T_table)
print_table("表2  30分钟内药材的水分浓度（kg/kg）", times_table, r_table_cm, C_table)

# ===================== 验证 =====================
print("\n===== 验证 =====")
print(f"t=1800s 中心温度梯度 T[1]-T[0] = {T[1]-T[0]:.3e}（应≈0）")
print(f"t=1800s 中心水分梯度 C[1]-C[0] = {C[1]-C[0]:.3e}（应≈0）")
print(f"t=1800s 温度径向: 中心 T[0]={T[0]:.4f} → 表面 T[N]={T[N]:.4f}（表面应更热）")
print(f"t=1800s 水分径向: 中心 C[0]={C[0]:.4f} → 表面 C[N]={C[N]:.4f}（表面应更干）")
print(f"t=1800s 烘房温度={T_air_interp(1800):.4f}℃, 烘房水分={C_air_interp(1800):.5f} kg/kg")

# ===================== 写 result1.xlsx =====================
def write_result1(path, times, r_vals, T_full, C_full):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "温度"
    ws.cell(1, 1, "时间\\到药材中心的距离")
    for j, rv in enumerate(r_vals):
        ws.cell(1, 2 + j, rv)
    for i, t in enumerate(times):
        ws.cell(2 + i, 1, t)
        for j in range(len(r_vals)):
            ws.cell(2 + i, 2 + j, round(float(T_full[i, j]), 4))

    ws2 = wb.create_sheet("水分浓度")
    ws2.cell(1, 1, "时间\\到药材中心的距离")
    for j, rv in enumerate(r_vals):
        ws2.cell(1, 2 + j, rv)
    for i, t in enumerate(times):
        ws2.cell(2 + i, 1, t)
        for j in range(len(r_vals)):
            ws2.cell(2 + i, 2 + j, round(float(C_full[i, j]), 4))

    wb.save(path)
    print(f"\n已写出 {path}（{len(times)} 行 × {len(r_vals)} 列 × 2 工作表）")

r_vals_round = [round(float(v) * 100, 1) for v in r_full]   # 0, 0.1, ..., 2.0（cm）
times_full = [(i + 1) * dt for i in range(n_steps)]
write_result1(os.path.join(RESULT_DIR, 'result1.xlsx'), times_full, r_vals_round, T_full, C_full)
