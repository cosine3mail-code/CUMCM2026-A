# -*- coding: utf-8 -*-
"""问题3物性参数与边界参数的单因素敏感性分析。

本脚本以同目录 ``问题3_求解器.py`` 的模型、网格、时间步、PCHIP
环境边界、后向 Euler、有限体积和联合 Picard 迭代为基准，不修改
原求解器。每次只扰动一个参数，默认扰动幅度为基准值的 ±5%。

物性参数 rho、cp、k、D 的“±5%”表示将完整经验函数整体乘以
0.95 或 1.05，不改变附录3经验式内部各项的相对结构。

边界温度的百分比以驱动力 Delta T_inf = T_inf - T0 为参数，避免
直接对摄氏温标做比例缩放；恒温段空气水分 C_inf 直接按比例扰动。
附件1覆盖的 0--4 h 实测边界始终保持不变，4 h之后的基准边界取
附件1最后两小时（7200--14400 s，含端点）观测值的算术平均。
"""

from __future__ import annotations

import argparse
import os
import time
from dataclasses import dataclass, replace

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.interpolate import PchipInterpolator


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
ATT1 = os.path.join(ROOT, "附件", "附件1.xlsx")
RESULT_DIR = os.path.join(SCRIPT_DIR, "结果", "敏感性分析")

R = 0.02
T0 = 28.0
C0 = 2.55
C_TARGET = 0.15
N = 200
DR = R / N
DT = 60.0
T_MAX = 500000.0
MAX_PICARD = 50
T_TOL = 1.0e-6
C_TOL = 1.0e-8
BISECTION_STEPS = 16


@dataclass(frozen=True)
class Result:
    drying_time_s: float
    center_moisture: float
    surface_moisture: float
    center_temperature: float
    surface_temperature: float
    mean_picard: float
    max_picard: int
    min_moisture: float
    max_temperature_overshoot: float
    elapsed_s: float


def read_attachment(path: str):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, max_col=3, values_only=True))
    finally:
        workbook.close()
    data = np.asarray(rows, dtype=float)
    return data[:, 0], data[:, 1], data[:, 2]


T_AIR_TIME, T_AIR_VALUE, C_AIR_VALUE = read_attachment(ATT1)
T_AIR = PchipInterpolator(T_AIR_TIME, T_AIR_VALUE)
C_AIR = PchipInterpolator(T_AIR_TIME, C_AIR_VALUE)
LAST_TWO_HOURS = (T_AIR_TIME >= 7200.0) & (T_AIR_TIME <= 14400.0)
if not np.any(LAST_TWO_HOURS):
    raise ValueError("附件1中没有7200--14400 s的观测值，无法计算最后两小时均值。")
BASE_T_INF = float(np.mean(T_AIR_VALUE[LAST_TWO_HOURS]))
BASE_C_INF = float(np.mean(C_AIR_VALUE[LAST_TWO_HOURS]))


@dataclass(frozen=True)
class Parameters:
    rho_scale: float = 1.0
    cp_scale: float = 1.0
    k_scale: float = 1.0
    D_scale: float = 1.0
    h: float = 25.0
    hm: float = 8.0e-7
    T_inf: float = BASE_T_INF
    C_inf: float = BASE_C_INF


def rho(C, p):
    return p.rho_scale * (650.0 + 128.0 * C)


def cp(C, p):
    return p.cp_scale * (1450.0 + 2736.0 * C / (C + 1.0))


def kcond(C, p):
    return p.k_scale * (0.21 + 0.38 * C / (C + 1.0))


def diffusivity(C, T_C, p):
    return p.D_scale * 2.4e-3 * np.exp(-0.45 / C) * np.exp(
        -3850.0 / (T_C + 273.15)
    )


def T_boundary(t, p):
    return float(T_AIR(t)) if t <= 14400.0 else p.T_inf


def C_boundary(t, p):
    return float(C_AIR(t)) if t <= 14400.0 else p.C_inf


def thomas(a, b, c, d):
    a = a.copy()
    b = b.copy()
    c = c.copy()
    d = d.copy()
    n = len(b)
    upper = np.zeros(n)
    rhs = np.zeros(n)
    upper[0] = c[0] / b[0]
    rhs[0] = d[0] / b[0]
    for i in range(1, n):
        pivot = b[i] - a[i] * upper[i - 1]
        if not np.isfinite(pivot) or abs(pivot) < 1.0e-30:
            raise ArithmeticError(f"Thomas 法在第 {i} 行遇到无效主元。")
        upper[i] = c[i] / pivot
        rhs[i] = (d[i] - a[i] * rhs[i - 1]) / pivot
    x = np.zeros(n)
    x[-1] = rhs[-1]
    for i in range(n - 2, -1, -1):
        x[i] = rhs[i] - upper[i] * x[i + 1]
    return x


def solve_heat(rhocp, k_face, T_old, t_new, dt_step, p):
    a = np.zeros(N + 1)
    b = np.zeros(N + 1)
    c = np.zeros(N + 1)
    d = np.zeros(N + 1)
    lam0 = k_face[0] * dt_step / (rhocp[0] * DR**2)
    b[0] = 1.0 + 4.0 * lam0
    c[0] = -4.0 * lam0
    d[0] = T_old[0]
    for i in range(1, N):
        lam_l = k_face[i - 1] * dt_step / (rhocp[i] * DR**2)
        lam_r = k_face[i] * dt_step / (rhocp[i] * DR**2)
        a[i] = -(i - 0.5) / i * lam_l
        c[i] = -(i + 0.5) / i * lam_r
        b[i] = 1.0 - a[i] - c[i]
        d[i] = T_old[i]
    ah = 2.0 * N * p.h * dt_step / ((N - 0.25) * rhocp[N] * DR)
    lam_n = k_face[N - 1] * dt_step / (rhocp[N] * DR**2)
    bh = 2.0 * (N - 0.5) * lam_n / (N - 0.25)
    a[N] = -bh
    b[N] = 1.0 + ah + bh
    d[N] = T_old[N] + ah * T_boundary(t_new, p)
    return thomas(a, b, c, d)


def solve_mass(D_face, C_old, t_new, dt_step, p):
    lam = D_face * dt_step / DR**2
    a = np.zeros(N + 1)
    b = np.zeros(N + 1)
    c = np.zeros(N + 1)
    d = np.zeros(N + 1)
    b[0] = 1.0 + 4.0 * lam[0]
    c[0] = -4.0 * lam[0]
    d[0] = C_old[0]
    for i in range(1, N):
        a[i] = -(i - 0.5) / i * lam[i - 1]
        c[i] = -(i + 0.5) / i * lam[i]
        b[i] = 1.0 - a[i] - c[i]
        d[i] = C_old[i]
    am = 2.0 * N * p.hm * dt_step / ((N - 0.25) * DR)
    bm = 2.0 * (N - 0.5) * lam[N - 1] / (N - 0.25)
    a[N] = -bm
    b[N] = 1.0 + am + bm
    d[N] = C_old[N] + am * C_boundary(t_new, p)
    return thomas(a, b, c, d)


def advance_one_step(T_old, C_old, t_new, dt_step, p):
    T_new = T_old.copy()
    C_new = C_old.copy()
    for iteration in range(1, MAX_PICARD + 1):
        T_previous = T_new.copy()
        C_previous = C_new.copy()
        rhocp = rho(C_previous, p) * cp(C_previous, p)
        k_face = kcond(0.5 * (C_previous[:-1] + C_previous[1:]), p)
        T_new = solve_heat(rhocp, k_face, T_old, t_new, dt_step, p)
        C_face = 0.5 * (C_previous[:-1] + C_previous[1:])
        T_face = 0.5 * (T_new[:-1] + T_new[1:])
        D_face = diffusivity(C_face, T_face, p)
        C_new = solve_mass(D_face, C_old, t_new, dt_step, p)
        err_t = float(np.max(np.abs(T_new - T_previous)))
        err_c = float(np.max(np.abs(C_new - C_previous)))
        if err_t < T_TOL and err_c < C_TOL:
            return T_new, C_new, iteration
    raise RuntimeError(
        f"t={t_new:.6f} s 的 Picard 迭代在 {MAX_PICARD} 次内未收敛；"
        f"残差为 ({err_t:.3e}, {err_c:.3e})。"
    )


def bisect_threshold(T_low, C_low, t_low, T_high, C_high, t_high, p):
    for _ in range(BISECTION_STEPS):
        t_mid = 0.5 * (t_low + t_high)
        T_mid, C_mid, _ = advance_one_step(
            T_low, C_low, t_mid, t_mid - t_low, p
        )
        if C_mid[0] < C_TARGET:
            t_high, T_high, C_high = t_mid, T_mid, C_mid
        else:
            t_low, T_low, C_low = t_mid, T_mid, C_mid
    return t_high, T_high, C_high


def drying_result(p):
    start = time.perf_counter()
    T = np.full(N + 1, T0)
    C = np.full(N + 1, C0)
    iterations = []
    min_moisture = C0
    max_overshoot = 0.0
    for n in range(int(T_MAX / DT)):
        t_new = (n + 1) * DT
        T_old = T.copy()
        C_old = C.copy()
        T, C, n_iter = advance_one_step(T_old, C_old, t_new, DT, p)
        iterations.append(n_iter)
        min_moisture = min(min_moisture, float(np.min(C)))
        ambient_max = max(T0, T_boundary(t_new, p))
        max_overshoot = max(max_overshoot, float(np.max(T) - ambient_max))
        if not np.all(np.isfinite(T)) or not np.all(np.isfinite(C)):
            raise ArithmeticError(f"t={t_new:g} s 出现非有限场值。")
        if float(np.min(C)) < -1.0e-12:
            raise ArithmeticError(f"t={t_new:g} s 出现负含水率。")
        if C[0] < C_TARGET:
            td, T_dry, C_dry = bisect_threshold(
                T_old, C_old, t_new - DT, T, C, t_new, p
            )
            return Result(
                drying_time_s=td,
                center_moisture=float(C_dry[0]),
                surface_moisture=float(C_dry[-1]),
                center_temperature=float(T_dry[0]),
                surface_temperature=float(T_dry[-1]),
                mean_picard=float(np.mean(iterations)),
                max_picard=int(np.max(iterations)),
                min_moisture=min_moisture,
                max_temperature_overshoot=max_overshoot,
                elapsed_s=time.perf_counter() - start,
            )
    raise RuntimeError(f"{T_MAX / 3600.0:.1f} h 内未达到烘干阈值。")


PARAMETER_DEFINITIONS = [
    ("rho", "密度函数 rho(C)", "物性", "rho_scale"),
    ("cp", "比热函数 cp(C)", "物性", "cp_scale"),
    ("k", "导热系数函数 k(C)", "物性", "k_scale"),
    ("D", "扩散系数函数 D(C,T)", "物性", "D_scale"),
    ("h", "表面对流换热系数 h", "边界", "h"),
    ("hm", "表面对流传质系数 hm", "边界", "hm"),
    ("DeltaT_inf", "恒温段温差 T_inf-T0", "边界", "delta_T_inf"),
    ("C_inf", "恒温段空气水分 C_inf", "边界", "C_inf"),
]


def perturbed(base, field, factor):
    if field == "delta_T_inf":
        return replace(base, T_inf=T0 + (base.T_inf - T0) * factor)
    value = getattr(base, field)
    return replace(base, **{field: value * factor})


def parameter_value(p, field):
    if field == "delta_T_inf":
        return p.T_inf - T0
    return getattr(p, field)


def write_results(rows, delta, baseline):
    os.makedirs(RESULT_DIR, exist_ok=True)
    output = os.path.join(RESULT_DIR, "问题3_物性与边界参数敏感性.xlsx")
    workbook = openpyxl.Workbook()
    detail = workbook.active
    detail.title = "单因素结果"
    headers = [
        "类别", "参数", "含义", "扰动", "参数值", "烘干时间/h",
        "相对基准变化/%", "中心终值C", "表面终值C", "中心终值T/℃",
        "表面终值T/℃", "平均Picard次数", "最大Picard次数",
        "全程最小C", "最大温度超调/℃", "计算耗时/s",
    ]
    detail.append(headers)
    for row in rows:
        detail.append(row)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    detail.freeze_panes = "A2"
    widths = [10, 15, 29, 10, 14, 15, 17, 13, 13, 15, 15, 17, 17, 13, 18, 13]
    for i, width in enumerate(widths, 1):
        detail.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    summary = workbook.create_sheet("敏感度汇总")
    summary_headers = [
        "类别", "参数", "含义", "-扰动时间/h", "基准时间/h", "+扰动时间/h",
        "-扰动相对变化/%", "+扰动相对变化/%", "归一化敏感度S",
        "绝对敏感度|S|", "变化方向是否符合预期",
    ]
    summary.append(summary_headers)
    baseline_h = baseline.drying_time_s / 3600.0
    grouped = {}
    for row in rows:
        if row[3] in (f"-{delta * 100:g}%", f"+{delta * 100:g}%"):
            grouped.setdefault(row[1], {})[row[3]] = row
    for key, label, category, _ in PARAMETER_DEFINITIONS:
        low = grouped[key][f"-{delta * 100:g}%"]
        high = grouped[key][f"+{delta * 100:g}%"]
        low_h = low[5]
        high_h = high[5]
        sensitivity = (high_h - low_h) / (2.0 * delta * baseline_h)
        expected = "是" if (
            (key in {"rho", "cp", "C_inf"} and sensitivity >= 0.0)
            or (key in {"k", "D", "h", "hm", "DeltaT_inf"} and sensitivity <= 0.0)
        ) else "需复核"
        summary.append([
            category, key, label, low_h, baseline_h, high_h,
            low[6], high[6], sensitivity, abs(sensitivity), expected,
        ])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    summary.freeze_panes = "A2"
    for i, width in enumerate([10, 15, 29, 16, 16, 16, 20, 20, 19, 16, 22], 1):
        summary.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    workbook.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description="问题3物性与边界参数敏感性分析")
    parser.add_argument(
        "--delta", type=float, default=0.05,
        help="单因素相对扰动幅度，默认0.05，即±5%%。",
    )
    args = parser.parse_args()
    if not 0.0 < args.delta < 0.5:
        parser.error("--delta 必须位于 (0, 0.5) 内。")

    delta = args.delta
    base = Parameters()
    print(f"问题3单因素敏感性分析：N={N}, dt={DT:g}s, 扰动=±{delta * 100:g}%")
    print("物性经验函数整体缩放；0--4 h附件1边界保持不变。")
    print(
        "4 h后采用附件1最后两小时平均值："
        f"T_inf={BASE_T_INF:.9f}℃，C_inf={BASE_C_INF:.9f} kg/kg。"
    )
    baseline = drying_result(base)
    baseline_h = baseline.drying_time_s / 3600.0
    print(f"基准：{baseline_h:.6f} h")

    rows = [[
        "基准", "baseline", "全部基准参数", "0%", 1.0, baseline_h, 0.0,
        baseline.center_moisture, baseline.surface_moisture,
        baseline.center_temperature, baseline.surface_temperature,
        baseline.mean_picard, baseline.max_picard, baseline.min_moisture,
        baseline.max_temperature_overshoot, baseline.elapsed_s,
    ]]
    for key, label, category, field in PARAMETER_DEFINITIONS:
        for sign, factor in (("-", 1.0 - delta), ("+", 1.0 + delta)):
            p = perturbed(base, field, factor)
            result = drying_result(p)
            td_h = result.drying_time_s / 3600.0
            relative = (td_h - baseline_h) / baseline_h * 100.0
            rows.append([
                category, key, label, f"{sign}{delta * 100:g}%",
                parameter_value(p, field), td_h, relative,
                result.center_moisture, result.surface_moisture,
                result.center_temperature, result.surface_temperature,
                result.mean_picard, result.max_picard, result.min_moisture,
                result.max_temperature_overshoot, result.elapsed_s,
            ])
            print(
                f"{key:>10s} {sign}{delta * 100:g}%：{td_h:.6f} h "
                f"({relative:+.4f}%)，Picard最大{result.max_picard}次"
            )

    output = write_results(rows, delta, baseline)
    print(f"结果已写入：{output}")


if __name__ == "__main__":
    main()
