# -*- coding: utf-8 -*-
"""
问题3 求解器：确定药材烘干所需时间（各处水分浓度 C < 0.15 kg/kg）
模型：同问题2，耦合传热传质、变物性(附录3)、一维径向圆柱
边界：预热段(0-4h 附件1) + 恒温干燥段(>4h, T=50℃, C=0.05；末值噪声敏感性见 敏感性扫描.py)
方法：后向 Euler 隐式 + 有限体积 + Thomas + 耦合 Picard 迭代
判据：C(r=0,t) < 0.15（水分径向分布中心最高，全剖面单调性见 单调性检查.py）
输出：烘干时长(h) + 表5(每6h) + result3.xlsx(每60s、每0.1cm)
"""
import os
import numpy as np
import openpyxl
from scipy.interpolate import PchipInterpolator

# ===================== 路径 =====================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
RESULT_DIR = os.path.join(ROOT, '问题3', '结果')
os.makedirs(RESULT_DIR, exist_ok=True)
ATT1 = os.path.join(ROOT, '附件', '附件1.xlsx')

# ===================== 参数 =====================
R    = 0.02
h    = 25.0
hm   = 8e-7
T0   = 28.0
C0   = 2.55
N    = 200
dr   = R / N
dt   = 60.0               # 时间步长 s（与 result3 输出间隔一致）
C_target = 0.15           # 烘干判据
t_max = 500000.0          # 求解上限 ~5.8 天（提前达标则 break）
n_max = int(t_max / dt)


def rho(C):   return 650.0 + 128.0 * C
def cp(C):    return 1450.0 + 2736.0 * C / (C + 1.0)
def kcond(C): return 0.21 + 0.38 * C / (C + 1.0)
def Dfun(C, T_C):
    return 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / (T_C + 273.15))


# ===================== 读附件1 =====================
def read_att1(path):
    wb = openpyxl.load_workbook(path, data_only=True); ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([r[0] for r in rows], float),
            np.array([r[1] for r in rows], float),
            np.array([r[2] for r in rows], float))

t_air, T_air, C_air = read_att1(ATT1)
T_air_interp = PchipInterpolator(t_air, T_air)
C_air_interp = PchipInterpolator(t_air, C_air)

def T_bc(t): return T_air_interp(t) if t <= 14400.0 else 49.97   # 平台均值（附件1 末2h平均）
def C_bc(t): return C_air_interp(t) if t <= 14400.0 else 0.0499


# ===================== Thomas =====================
def thomas(a, b, c, d):
    a = a.copy(); b = b.copy(); c = c.copy(); d = d.copy(); n = len(b)
    cp_ = np.zeros(n); dp_ = np.zeros(n)
    cp_[0] = c[0] / b[0]; dp_[0] = d[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp_[i - 1]
        cp_[i] = c[i] / m
        dp_[i] = (d[i] - a[i] * dp_[i - 1]) / m
    x = np.zeros(n); x[-1] = dp_[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp_[i] - cp_[i] * x[i + 1]
    return x


# ===================== 求解器（复用问题2 的耦合 Picard） =====================
def solve_heat(rhocp, k_face, T_old, tn):
    aT = np.zeros(N + 1); bT = np.zeros(N + 1); cT = np.zeros(N + 1); dT = np.zeros(N + 1)
    lam0 = k_face[0] * dt / (rhocp[0] * dr**2)
    bT[0] = 1 + 4 * lam0; cT[0] = -4 * lam0; dT[0] = T_old[0]
    for i in range(1, N):
        lamL = k_face[i - 1] * dt / (rhocp[i] * dr**2)
        lamR = k_face[i] * dt / (rhocp[i] * dr**2)
        aT[i] = -(i - 0.5) / i * lamL
        cT[i] = -(i + 0.5) / i * lamR
        bT[i] = 1 + (i - 0.5) / i * lamL + (i + 0.5) / i * lamR
        dT[i] = T_old[i]
    Ah = 2 * N * h * dt / ((N - 0.25) * rhocp[N] * dr)
    lamN = k_face[N - 1] * dt / (rhocp[N] * dr**2)
    Bh = 2 * (N - 0.5) * lamN / (N - 0.25)
    aT[N] = -Bh; bT[N] = 1 + Ah + Bh; dT[N] = T_old[N] + Ah * T_bc(tn)
    return thomas(aT, bT, cT, dT)


def solve_mass(D_face, C_old, tn):
    lam = D_face * dt / dr**2
    aC = np.zeros(N + 1); bC = np.zeros(N + 1); cC = np.zeros(N + 1); dC = np.zeros(N + 1)
    bC[0] = 1 + 4 * lam[0]; cC[0] = -4 * lam[0]; dC[0] = C_old[0]
    for i in range(1, N):
        aC[i] = -(i - 0.5) / i * lam[i - 1]
        cC[i] = -(i + 0.5) / i * lam[i]
        bC[i] = 1 + (i - 0.5) / i * lam[i - 1] + (i + 0.5) / i * lam[i]
        dC[i] = C_old[i]
    Am = 2 * N * hm * dt / ((N - 0.25) * dr)
    Bm = 2 * (N - 0.5) * lam[N - 1] / (N - 0.25)
    aC[N] = -Bm; bC[N] = 1 + Am + Bm; dC[N] = C_old[N] + Am * C_bc(tn)
    return thomas(aC, bC, cC, dC)


# ===================== 时间推进 =====================
T = np.full(N + 1, T0)
C = np.full(N + 1, C0)
idx_full = np.arange(0, N + 1, 10)          # 0..2cm 每0.1cm（21点）
idx_table = [0, 50, 100, 150, 200]          # 0,0.5,1,1.5,2 cm

C_full = np.zeros((n_max, len(idx_full)))   # 每60s、每0.1cm
T_full = np.zeros((n_max, len(idx_full)))

n_done = 0
t_dry = None
iter_hist = []
for n in range(n_max):
    tn = (n + 1) * dt
    T_new = T.copy(); C_new = C.copy()
    for it in range(50):
        Tp = T_new.copy(); Cp = C_new.copy()
        rhocp = rho(Cp) * cp(Cp)
        k_face = kcond(0.5 * (Cp[:-1] + Cp[1:]))
        T_new = solve_heat(rhocp, k_face, T, tn)
        Cf = 0.5 * (Cp[:-1] + Cp[1:])
        Tf = 0.5 * (T_new[:-1] + T_new[1:])
        D_face = Dfun(Cf, Tf)
        C_new = solve_mass(D_face, C, tn)
        if np.max(np.abs(T_new - Tp)) < 1e-6 and np.max(np.abs(C_new - Cp)) < 1e-8:
            break
    T = T_new; C = C_new
    iter_hist.append(it + 1)
    C_full[n] = C[idx_full]; T_full[n] = T[idx_full]
    n_done = n + 1
    # 烘干判据：中心水分 < 0.15（首次满足）
    if C[0] < C_target:
        # 线性插值精确定位跨越时刻
        C_prev = C_full[n - 1, 0] if n > 0 else C0
        t_dry = (n + (C_target - C_prev) / (C[0] - C_prev)) * dt if n > 0 else dt
        break

if t_dry is None:
    print(f"警告：{t_max/3600:.1f} h 内未达到 C<0.15，中心 C={C[0]:.4f}")
    t_dry = n_done * dt

t_dry_h = t_dry / 3600.0
print(f"===== 问题3 结果 =====")
print(f"烘干所需时间：{t_dry:.1f} s = {t_dry_h:.4f} h")
print(f"结束时刻中心水分 C(r=0) = {C[0]:.4f} kg/kg（判据 <0.15）")
print(f"结束时刻表面水分 C(r=R) = {C[N]:.4f} kg/kg")
print(f"结束时刻温度：中心 {T[0]:.2f}℃ → 表面 {T[N]:.2f}℃")
print(f"Picard 迭代：平均 {np.mean(iter_hist):.1f} 次/步，最大 {np.max(iter_hist)} 次")

# ===================== 表5（每6h） =====================
r_table_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
times_6h = list(range(6 * 3600, int(t_dry) + 1, 6 * 3600))   # 6,12,18,...h 至烘干结束
# 确保最后一行是烘干结束时刻
times_6h = [t for t in times_6h if t <= t_dry]
if times_6h[-1] < t_dry - 1:
    times_6h.append(int(round(t_dry)))

print(f"\n表5  药材烘干过程的水分浓度（kg/kg）")
header = "时间/h | " + " | ".join(f"{v:g}cm" for v in r_table_cm)
print(header); print("-" * len(header))
C_table_rows = []
for ts in times_6h:
    idx = min(int(round(ts / dt)) - 1, n_done - 1)
    row = C_full[idx][[0, 5, 10, 15, 20]]   # idx_table 对应 idx_full 的 0,5,10,15,20
    C_table_rows.append(row)
    line = f"{ts/3600:>6.1f} | " + " | ".join(f"{v:.4f}" for v in row)
    print(line)

# ===================== 写 result3.xlsx =====================
r_vals = [round(float(v) * 100, 1) for v in (idx_full * dr)]
times = list(range(60, int(n_done * dt) + 1, 60))   # 60,120,... 至烘干结束
n_out = len(times)

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "水分浓度"
ws.cell(1, 1, "时间\\到药材中心的距离")
for j, rv in enumerate(r_vals):
    ws.cell(1, 2 + j, rv)
for i, t in enumerate(times):
    ws.cell(2 + i, 1, t)
    for j in range(len(r_vals)):
        ws.cell(2 + i, 2 + j, round(float(C_full[i, j]), 4)).number_format = '0.0000'
wb.save(os.path.join(RESULT_DIR, 'result3.xlsx'))

# 表5 存档
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "表5水分"
ws.cell(1, 1, "时间/h"); [ws.cell(1, 2 + j, f"{v:g}") for j, v in enumerate(r_table_cm)]
for i, ts in enumerate(times_6h):
    ws.cell(2 + i, 1, round(ts / 3600.0, 2))
    for j in range(5):
        ws.cell(2 + i, 2 + j, round(float(C_table_rows[i][j]), 4)).number_format = '0.0000'
wb.save(os.path.join(RESULT_DIR, '表5.xlsx'))

print(f"\n已写出 问题3/结果/result3.xlsx（{n_out} 行）与 表5.xlsx")
