#!/usr/bin/env python3
"""Plot the Question 3 moisture field as a space-time heatmap.

The source workbook is read-only.  The figure contains moisture only; the
temperature field is intentionally not plotted.  The target moisture
threshold C = 0.15 kg/kg is drawn as a highlighted red contour.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import colors
import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
INPUT_XLSX = PROJECT_ROOT / "计算结果" / "result3_计算结果.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT_PNG = OUTPUT_DIR / "问题三含水率时空热力图.png"
OUTPUT_PDF = OUTPUT_DIR / "问题三含水率时空热力图.pdf"


def read_moisture_field(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Read time, radius and moisture values from result3.xlsx."""
    if not path.is_file():
        raise FileNotFoundError(f"找不到输入文件：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < 2 or len(rows[0]) < 2:
        raise ValueError("问题三结果工作簿没有足够的时空数据。")

    radius_cm = np.asarray(rows[0][1:], dtype=float)
    time_s = np.asarray([row[0] for row in rows[1:]], dtype=float)
    moisture = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    if moisture.shape != (time_s.size, radius_cm.size):
        raise ValueError("问题三结果工作簿的时间/径向数据维度不一致。")
    if not np.all(np.isfinite(moisture)):
        raise ValueError("问题三含水率数据包含 NaN 或无穷值。")
    if np.any(np.diff(time_s) <= 0.0) or np.any(np.diff(radius_cm) <= 0.0):
        raise ValueError("时间和径向坐标必须严格递增。")
    return time_s, radius_cm, moisture


def make_time_ticks(end_time_h: float) -> np.ndarray:
    """Choose readable hour ticks for the complete Question 3 duration."""
    if end_time_h <= 12.0:
        interval = 2.0
    elif end_time_h <= 36.0:
        interval = 6.0
    elif end_time_h <= 72.0:
        interval = 12.0
    else:
        interval = 24.0
    ticks = np.arange(0.0, end_time_h + 0.5 * interval, interval)
    # Do not let a tick beyond the last calculated time expand the axis and
    # create an empty white strip at the right-hand side of the heatmap.
    ticks = ticks[ticks <= end_time_h + 1.0e-12]
    if ticks.size == 0 or not np.isclose(ticks[-1], end_time_h):
        ticks = np.append(ticks, end_time_h)
    return ticks


def cell_edges(coordinates: np.ndarray, lower: float, upper: float) -> np.ndarray:
    """Build finite-volume-like cell edges from nodal coordinates."""
    if coordinates.size < 2:
        return np.array([lower, upper], dtype=float)
    interior = 0.5 * (coordinates[:-1] + coordinates[1:])
    return np.concatenate(([lower], interior, [upper])).astype(float)


def interpolate_radially(
    radius_cm: np.ndarray,
    moisture: np.ndarray,
    display_node_count: int = 401,
) -> tuple[np.ndarray, np.ndarray]:
    """Linearly interpolate radial nodes only for smooth colour rendering.

    The calculation results and contour values are not changed.  This only
    removes the visible 0.1 cm-wide colour bands caused by flat plotting of
    the 21 output nodes in the workbook.
    """
    display_radius_cm = np.linspace(
        float(radius_cm[0]), float(radius_cm[-1]), display_node_count
    )
    display_moisture = np.vstack(
        [
            np.interp(display_radius_cm, radius_cm, moisture_at_time)
            for moisture_at_time in moisture
        ]
    )
    return display_radius_cm, display_moisture


def make_moisture_norm(moisture: np.ndarray) -> colors.FuncNorm:
    """Expand the colour resolution in the important 0.15--0.50 interval.

    The colour sequence matches the original moisture-concentration heatmap
    (``cividis_r``): low moisture is bright yellow and high moisture is dark
    blue.  A piecewise-linear map allocates more of the colourbar to the
    low-moisture interval without changing any plotted data values.
    """
    data_min = float(np.min(moisture))
    data_max = float(np.max(moisture))
    if data_max <= data_min:
        raise ValueError("含水率数据范围必须为正。")

    threshold = 0.15
    focus_end = 0.50
    # Keep the mapping valid even if a future result has a narrower range.
    x = np.array(
        [data_min, max(data_min, threshold), max(data_min, focus_end), data_max],
        dtype=float,
    )
    x = np.maximum.accumulate(x)
    x = np.unique(x)
    if x.size < 2:
        return colors.Normalize(vmin=data_min, vmax=data_max)  # type: ignore[return-value]

    y_template = np.array([0.0, 0.34, 0.68, 1.0], dtype=float)
    y = y_template[np.linspace(0, 3, x.size).round().astype(int)]

    def forward(value):
        return np.interp(value, x, y)

    def inverse(value):
        return np.interp(value, y, x)

    return colors.FuncNorm((forward, inverse), vmin=data_min, vmax=data_max)


def main() -> None:
    time_s, radius_cm, moisture = read_moisture_field(INPUT_XLSX)
    time_h = time_s / 3600.0
    end_time_h = float(time_h[-1])
    # Smoothly interpolate between the 21 radial output nodes for display.
    # Use cell edges so the plotted domain remains exactly [0, t_end] x
    # [0, 2 cm] without extrapolated or empty boundary strips.
    display_radius_cm, display_moisture = interpolate_radially(radius_cm, moisture)
    time_edges = cell_edges(time_h, 0.0, end_time_h)
    radius_edges = cell_edges(
        display_radius_cm,
        float(display_radius_cm.min()),
        float(display_radius_cm.max()),
    )
    time_grid, radius_grid = np.meshgrid(time_edges, radius_edges)
    time_nodes, radius_nodes = np.meshgrid(time_h, radius_cm)

    plt.rcParams.update(
        {
            "font.sans-serif": [
                "Microsoft YaHei",
                "SimHei",
                "Noto Sans CJK SC",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
            "font.size": 11,
        }
    )

    fig, ax = plt.subplots(figsize=(12.2, 5.8), constrained_layout=True)
    moisture_norm = make_moisture_norm(moisture)
    image = ax.pcolormesh(
        time_grid,
        radius_grid,
        display_moisture.T,
        shading="flat",
        cmap="cividis_r",
        norm=moisture_norm,
        rasterized=True,
    )

    # Regular contours provide spatial context.  The 0.15 threshold is
    # plotted separately so it remains visually dominant even near the lower
    # end of the moisture range.
    regular_levels = np.array([0.30, 0.60, 0.90, 1.20, 1.50, 1.80, 2.10, 2.40])
    regular_levels = regular_levels[
        (regular_levels > float(moisture.min()))
        & (regular_levels < float(moisture.max()))
    ]
    if regular_levels.size:
        regular = ax.contour(
            time_nodes,
            radius_nodes,
            moisture.T,
            levels=regular_levels,
            colors="white",
            linewidths=0.85,
            alpha=0.86,
        )
        regular_labels = ax.clabel(
            regular,
            inline=True,
            fontsize=8,
            fmt="%.2f",
            colors="white",
        )
        for label in regular_labels:
            label.set_fontweight("bold")

    threshold = 0.15
    if float(moisture.min()) <= threshold <= float(moisture.max()):
        threshold_contour = ax.contour(
            time_nodes,
            radius_nodes,
            moisture.T,
            levels=[threshold],
            colors="#d62728",
            linewidths=2.4,
            alpha=1.0,
        )
        threshold_labels = ax.clabel(
            threshold_contour,
            inline=True,
            fontsize=9,
            fmt={threshold: "0.15"},
            colors="#d62728",
        )
        for label in threshold_labels:
            label.set_fontweight("bold")
            label.set_bbox(
                dict(facecolor="white", edgecolor="#d62728", alpha=0.82, pad=1.5)
            )

    ax.set_xlabel("时间 / h")
    ax.set_ylabel("到药材中心的距离 / cm")
    ax.set_xlim(0.0, end_time_h)
    ax.set_ylim(float(radius_cm.min()), float(radius_cm.max()))
    ax.set_xticks(make_time_ticks(end_time_h))
    ax.set_yticks(np.arange(0.0, 2.01, 0.5))
    ax.tick_params(direction="in", length=4)
    ax.grid(False)

    colorbar = fig.colorbar(image, ax=ax, orientation="vertical", pad=0.025)
    colorbar.set_label("干基含水率 / (kg/kg)", labelpad=8)
    colorbar.set_ticks([0.15, 0.20, 0.30, 0.40, 0.50, 1.00, 1.50, 2.00, 2.50])
    colorbar.ax.tick_params(length=3)
    for tick, value in zip(colorbar.ax.get_yticklabels(), colorbar.get_ticks()):
        if np.isclose(value, threshold):
            tick.set_color("#d62728")
            tick.set_fontweight("bold")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUTPUT_PNG, dpi=300, bbox_inches="tight")
    fig.savefig(OUTPUT_PDF, bbox_inches="tight")
    plt.close(fig)
    print(f"问题三含水率热力图 PNG 已保存至：{OUTPUT_PNG}")
    print(f"问题三含水率热力图 PDF 已保存至：{OUTPUT_PDF}")


if __name__ == "__main__":
    main()
