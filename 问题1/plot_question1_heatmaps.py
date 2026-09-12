#!/usr/bin/env python3
"""Plot temperature and moisture space-time heatmaps from a result workbook."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
INPUT_XLSX = PROJECT_ROOT / "计算结果" / "result2_计算结果_PCHIP.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT_PNG = OUTPUT_DIR / "问题二温度与含水率时空热力图_等温线均匀.png"
OUTPUT_PDF = OUTPUT_DIR / "问题二温度与含水率时空热力图_等温线均匀.pdf"


def read_field(sheet) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError(f"工作表 {sheet.title!r} 没有足够数据。")

    radius_cm = np.asarray(rows[0][1:], dtype=float)
    time_s = np.asarray([row[0] for row in rows[1:]], dtype=float)
    values = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    if values.shape != (time_s.size, radius_cm.size):
        raise ValueError(f"工作表 {sheet.title!r} 的数据维度不一致。")
    return time_s, radius_cm, values


def make_time_ticks(end_time_min: float) -> np.ndarray:
    """Return readable minute ticks covering the complete result duration."""
    if end_time_min <= 30.0:
        interval = 5.0
    elif end_time_min <= 60.0:
        interval = 10.0
    elif end_time_min <= 180.0:
        interval = 30.0
    elif end_time_min <= 360.0:
        interval = 60.0
    else:
        interval = 120.0
    return np.arange(0.0, end_time_min + 0.5 * interval, interval)


def make_time_balanced_isotherms(
    time_min: np.ndarray,
    temperature: np.ndarray,
    level_count: int = 9,
) -> np.ndarray:
    """Choose isotherms that cross the centreline at even time intervals."""
    target_times = np.linspace(0.0, float(time_min.max()), level_count + 2)[1:-1]
    centre_temperature = temperature[:, 0]
    levels = np.interp(target_times, time_min, centre_temperature)
    levels = np.unique(levels)
    return levels[
        (levels > float(temperature.min()))
        & (levels < float(temperature.max()))
    ]


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"找不到输入文件：{INPUT_XLSX}")

    workbook = load_workbook(INPUT_XLSX, read_only=True, data_only=True)
    if len(workbook.sheetnames) < 2:
        raise ValueError("输入工作簿至少需要温度和水分浓度两个工作表。")

    temperature_time, radius_cm, temperature = read_field(workbook[workbook.sheetnames[0]])
    moisture_time, moisture_radius_cm, moisture = read_field(workbook[workbook.sheetnames[1]])
    workbook.close()

    if not np.array_equal(temperature_time, moisture_time) or not np.array_equal(
        radius_cm, moisture_radius_cm
    ):
        raise ValueError("温度和水分浓度工作表的时间/径向坐标不一致。")

    # The workbook stores time in seconds and radius in centimetres.
    time_min = temperature_time / 60.0
    end_time_min = float(time_min.max())
    time_ticks = make_time_ticks(end_time_min)
    extent = (0.0, end_time_min, float(radius_cm.min()), float(radius_cm.max()))

    plt.rcParams.update(
        {
            "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "font.size": 11,
        }
    )
    fig, axes = plt.subplots(1, 2, figsize=(14.2, 6.0), constrained_layout=True)

    temperature_contours = make_time_balanced_isotherms(time_min, temperature)
    moisture_contours = np.linspace(float(moisture.min()), float(moisture.max()), 9)
    panels = (
        (
            axes[0],
            temperature,
            "(a) 温度场",
            "温度 / °C",
            "magma",
            temperature_contours,
        ),
        (
            axes[1],
            moisture,
            "(b) 含水率场",
            "干基含水率 / (kg/kg)",
            "cividis_r",
            moisture_contours,
        ),
    )
    for ax, field, title, colorbar_label, cmap, contour_levels in panels:
        # Excel stores [time, radius]; plotting coordinates are [x=time, y=radius].
        time_grid, radius_grid = np.meshgrid(time_min, radius_cm)
        image = ax.pcolormesh(
            time_grid,
            radius_grid,
            field.T,
            shading="auto",
            cmap=cmap,
        )
        contours = ax.contour(
            time_grid,
            radius_grid,
            field.T,
            levels=contour_levels,
            colors="white",
            linewidths=0.35,
            alpha=0.55,
        )
        labels = ax.clabel(
            contours,
            inline=True,
            fontsize=7,
            fmt="%.2f",
            colors="white",
        )
        
        ax.set_title(title, pad=10)
        ax.set_xlabel("时间 / min")
        ax.set_ylabel("到药材中心的距离 / cm")
        ax.set_xlim(extent[0], extent[1])
        ax.set_ylim(extent[2], extent[3])
        ax.set_xticks(time_ticks)
        ax.set_yticks(np.arange(0, 2.01, 0.5))
        ax.tick_params(direction="in", length=4)
        ax.grid(True, color="white", linewidth=0.35, alpha=0.18)
        colorbar = fig.colorbar(
            image,
            ax=ax,
            orientation="horizontal",
            pad=0.12,
            fraction=0.08,
            aspect=35,
        )
        colorbar.set_label(colorbar_label, labelpad=6)
        colorbar.ax.tick_params(length=3)

    fig.savefig(OUTPUT_PNG, dpi=300, bbox_inches="tight")
    fig.savefig(OUTPUT_PDF, bbox_inches="tight")
    plt.close(fig)
    print(f"热力图 PNG 已保存至：{OUTPUT_PNG}")
    print(f"热力图 PDF 已保存至：{OUTPUT_PDF}")


if __name__ == "__main__":
    main()
