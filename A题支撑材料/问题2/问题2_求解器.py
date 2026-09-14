# -*- coding: utf-8 -*-
"""问题2全过程耦合传热传质求解器。

模型：耦合传热传质，变物性(附录3)，一维径向圆柱。
  ρ  = 650 + 128C
  cp = 1450 + 2736·C/(C+1)
  k  = 0.21 + 0.38·C/(C+1)
  D  = 2.4e-3·e^(-0.45/C)·e^(-3850/T)，T 单位 K
  h=25、hm=8e-7 沿用附录2
方法：后向 Euler 隐式 + 有限体积 + Thomas 三对角 + 耦合 Picard 迭代
输入：``../附件/附件1.xlsx``；输出：本目录 ``结果/``。
运行方式：在项目根目录执行 ``python 问题2/问题2_求解器.py``。
"""
import os
import numpy as np
import openpyxl
from scipy.interpolate import PchipInterpolator

# ===================== 路径 =====================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)          # A题/
RESULT_DIR = os.path.join(ROOT, '问题2', '结果')
os.makedirs(RESULT_DIR, exist_ok=True)
ATT1 = os.path.join(ROOT, '附件', '附件1.xlsx')

# ===================== 参数 =====================
R    = 0.02          # 半径 m
h    = 25.0          # 对流换热系数 W/(m^2·K)
hm   = 8e-7          # 对流传质系数 m/s
T0   = 28.0          # 初始温度 ℃
C0   = 2.55          # 初始水分浓度 kg/kg
N    = 200           # 径向网格数
dr   = R / N
t_end = 10800.0      # 报告 3 h（3600*3）
dt   = 1.0           # 时间步长 s
n_steps = int(round(t_end / dt))


# ===================== 附录3 物性 =====================
def rho(C):    return 650.0 + 128.0 * C
def cp(C):     return 1450.0 + 2736.0 * C / (C + 1.0)
def kcond(C):  return 0.21 + 0.38 * C / (C + 1.0)
def Dfun(C, T_C):
    """水分浓度扩散系数，T_C 为摄氏温度（内部转开尔文）"""
    return 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / (T_C + 273.15))


# ===================== 读附件1 =====================
def read_att1(path):
    wb = openpyxl.load_workbook(path, data_only=True); ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]
    return (np.array([r[0] for r in rows], float),
            np.array([r[1] for r in rows], float),
            np.array([r[2] for r in rows], float))

t_air, T_air, C_air = read_att1(ATT1)
T_air_interp = PchipInterpolator(t_air, T_air)
C_air_interp = PchipInterpolator(t_air, C_air)


# ===================== 环境边界（预热段附件1 + 恒温段常数） =====================
def T_bc(t):
    return T_air_interp(t) if t <= 14400.0 else 49.97   # 平台均值（附件1 末2h平均）

def C_bc(t):
    return C_air_interp(t) if t <= 14400.0 else 0.0499


# ===================== Thomas =====================
def thomas(a, b, c, d):
    a = a.copy(); b = b.copy(); c = c.copy(); d = d.copy(); n = len(b)
    cp_ = np.zeros(n); dp_ = np.zeros(n)
    cp_[0] = c[0] / b[0]; dp_[0] = d[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp_[i - 1]
        cp_[i] = c[i] / m
        dp_[i] = (d[i] - a[i] * dp_[i - 1]) / m
    x = np.zeros(n); x[-1] = dp_[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp_[i] - cp_[i] * x[i + 1]
    return x


# ===================== 热方程(变物性，系数由 C 决定) =====================
def solve_heat(rhocp, k_face, T_old, tn):
    aT = np.zeros(N + 1); bT = np.zeros(N + 1); cT = np.zeros(N + 1); dT = np.zeros(N + 1)
    lam0 = k_face[0] * dt / (rhocp[0] * dr**2)                 # 中心
    bT[0] = 1 + 4 * lam0; cT[0] = -4 * lam0; dT[0] = T_old[0]
    for i in range(1, N):
        lamL = k_face[i - 1] * dt / (rhocp[i] * dr**2)
        lamR = k_face[i] * dt / (rhocp[i] * dr**2)
        aT[i] = -(i - 0.5) / i * lamL
        cT[i] = -(i + 0.5) / i * lamR
        bT[i] = 1 + (i - 0.5) / i * lamL + (i + 0.5) / i * lamR
        dT[i] = T_old[i]
    Ah = 2 * N * h * dt / ((N - 0.25) * rhocp[N] * dr)          # 表面(半格点)
    lamN = k_face[N - 1] * dt / (rhocp[N] * dr**2)
    Bh = 2 * (N - 0.5) * lamN / (N - 0.25)
    aT[N] = -Bh; bT[N] = 1 + Ah + Bh; dT[N] = T_old[N] + Ah * T_bc(tn)
    return thomas(aT, bT, cT, dT)


# ===================== 水分方程（变物性，D 由 C、T 决定） =====================
def solve_mass(D_face, C_old, tn):
    lam = D_face * dt / dr**2
    aC = np.zeros(N + 1); bC = np.zeros(N + 1); cC = np.zeros(N + 1); dC = np.zeros(N + 1)
    bC[0] = 1 + 4 * lam[0]; cC[0] = -4 * lam[0]; dC[0] = C_old[0]   # 中心
    for i in range(1, N):
        aC[i] = -(i - 0.5) / i * lam[i - 1]
        cC[i] = -(i + 0.5) / i * lam[i]
        bC[i] = 1 + (i - 0.5) / i * lam[i - 1] + (i + 0.5) / i * lam[i]
        dC[i] = C_old[i]
    Am = 2 * N * hm * dt / ((N - 0.25) * dr)                      # 表面(半格点)
    Bm = 2 * (N - 0.5) * lam[N - 1] / (N - 0.25)
    aC[N] = -Bm; bC[N] = 1 + Am + Bm; dC[N] = C_old[N] + Am * C_bc(tn)
    return thomas(aC, bC, cC, dC)


# ===================== 时间推进（耦合 Picard 迭代） =====================
T = np.full(N + 1, T0)
C = np.full(N + 1, C0)

idx_full = np.arange(0, N + 1, 10)          # 0, 0.1, ..., 2.0 cm（21点）
idx_table = [0, 50, 100, 150, 200]          # 0, 0.5, 1, 1.5, 2 cm
r_full = idx_full * dr
r_table = np.array(idx_table) * dr

T_full = np.zeros((n_steps, len(idx_full)))
C_full = np.zeros((n_steps, len(idx_full)))

times_table = [1800, 3600, 5400, 7200, 9000, 10800]   # 0.5,1,1.5,2,2.5,3 h
tm = {t: i for i, t in enumerate(times_table)}
T_table = np.zeros((len(times_table), len(idx_table)))
C_table = np.zeros((len(times_table), len(idx_table)))

iter_hist = []

for n in range(n_steps):
    tn = (n + 1) * dt
    T_new = T.copy(); C_new = C.copy()
    for it in range(50):
        Tp = T_new.copy(); Cp = C_new.copy()
        # 物性（用上一迭代值）
        rhocp = rho(Cp) * cp(Cp)
        k_face = kcond(0.5 * (Cp[:-1] + Cp[1:]))
        # 热方程（线性，系数冻结）→ T_new
        T_new = solve_heat(rhocp, k_face, T, tn)
        # 水分方程采用当前浓度与新温度计算 D，并按 Gauss-Seidel 次序更新。
        Cf = 0.5 * (Cp[:-1] + Cp[1:])
        Tf = 0.5 * (T_new[:-1] + T_new[1:])
        D_face = Dfun(Cf, Tf)
        C_new = solve_mass(D_face, C, tn)
        # 收敛判据
        if np.max(np.abs(T_new - Tp)) < 1e-6 and np.max(np.abs(C_new - Cp)) < 1e-8:
            break
    T = T_new; C = C_new
    iter_hist.append(it)
    T_full[n] = T[idx_full]; C_full[n] = C[idx_full]
    if tn in tm:
        j = tm[tn]; T_table[j] = T[idx_table]; C_table[j] = C[idx_table]

# ===================== 输出表3/表4 =====================
def print_table(title, times_h, r_cm, data):
    print(f"\n{title}")
    header = "时间/h | " + " | ".join(f"{v:g}cm" for v in r_cm)
    print(header); print("-" * len(header))
    for i, th in enumerate(times_h):
        line = f"{th:>6.1f} | " + " | ".join(f"{data[i, j]:.4f}" for j in range(len(r_cm)))
        print(line)

times_h = [t / 3600.0 for t in times_table]
r_table_cm = r_table * 100
print_table("表3  3小时内药材的温度（℃）", times_h, r_table_cm, T_table)
print_table("表4  3小时内药材的水分浓度（kg/kg）", times_h, r_table_cm, C_table)

# ===================== 验证 =====================
print("\n===== 验证 =====")
print(f"t=3h 中心温度梯度 T[1]-T[0] = {T[1]-T[0]:.3e}（应≈0）")
print(f"t=3h 中心水分梯度 C[1]-C[0] = {C[1]-C[0]:.3e}（应≈0）")
print(f"t=3h 温度径向: 中心 {T[0]:.4f} → 表面 {T[N]:.4f}（表面应更热）")
print(f"t=3h 水分径向: 中心 {C[0]:.4f} → 表面 {C[N]:.4f}（表面应更干）")
print(f"Picard 迭代次数: 平均 {np.mean(iter_hist):.1f}，最大 {np.max(iter_hist)}")
print(f"t=3h 烘房温度 {T_bc(10800):.4f}℃, 烘房水分 {C_bc(10800):.5f} kg/kg")

# 质量守恒（总含水量单调下降）
r = np.arange(N + 1) * dr
print(f"总含水量比(3h/初始) = {np.trapezoid(C * r, r) / np.trapezoid(C0 * r, r):.6f}（应<1）")

# ===================== 写 result2.xlsx =====================
r_vals = [round(float(v) * 100, 1) for v in r_full]
times = list(range(1, n_steps + 1))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
ws.cell(1, 1, "时间\\到药材中心的距离")
for j, rv in enumerate(r_vals):
    ws.cell(1, 2 + j, rv)
for i, t in enumerate(times):
    ws.cell(2 + i, 1, t)
    for j in range(len(r_vals)):
        ws.cell(2 + i, 2 + j, round(float(T_full[i, j]), 4)).number_format = '0.0000'
ws2 = wb.create_sheet("水分浓度")
ws2.cell(1, 1, "时间\\到药材中心的距离")
for j, rv in enumerate(r_vals):
    ws2.cell(1, 2 + j, rv)
for i, t in enumerate(times):
    ws2.cell(2 + i, 1, t)
    for j in range(len(r_vals)):
        ws2.cell(2 + i, 2 + j, round(float(C_full[i, j]), 4)).number_format = '0.0000'
wb.save(os.path.join(RESULT_DIR, 'result2.xlsx'))

# 表3/表4
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "表3温度"
ws.cell(1, 1, "时间/h"); [ws.cell(1, 2 + j, f"{v:g}") for j, v in enumerate(r_table_cm)]
for i, th in enumerate(times_h):
    ws.cell(2 + i, 1, round(th, 1))
    for j in range(5):
        ws.cell(2 + i, 2 + j, round(float(T_table[i, j]), 4)).number_format = '0.0000'
ws2 = wb.create_sheet("表4水分")
ws2.cell(1, 1, "时间/h"); [ws2.cell(1, 2 + j, f"{v:g}") for j, v in enumerate(r_table_cm)]
for i, th in enumerate(times_h):
    ws2.cell(2 + i, 1, round(th, 1))
    for j in range(5):
        ws2.cell(2 + i, 2 + j, round(float(C_table[i, j]), 4)).number_format = '0.0000'
wb.save(os.path.join(RESULT_DIR, '表3表4.xlsx'))

print(f"\n已写出 问题2/结果/result2.xlsx 与 表3表4.xlsx")
