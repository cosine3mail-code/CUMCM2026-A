#!/usr/bin/env python3
"""复核问题1的一维径向假设并绘制图8。

脚本采用相同初值、附件2物性和附件1分段线性边界，分别推进：
1. 仅在圆柱侧面换热、传质的一维径向模型；
2. 侧面与物理端面均换热、传质的半长 ``r-z`` 二维轴对称模型。

两模型的差异仅来自圆形端面的处理。计算完成后，脚本比较端面影响长度、
中面径向分布和轴向平均结果，并将图8写入本目录的 ``图`` 文件夹。
附件原始数据、问题1正式结果工作簿及其他定稿文件均不会被修改。
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.sparse import coo_matrix, csc_matrix
from scipy.sparse.linalg import factorized, spsolve


@dataclass(frozen=True)
class Parameters:
    """问题1几何、初值与附件2物性参数。"""

    length: float = 0.25
    radius: float = 0.02
    initial_temperature: float = 28.0
    initial_moisture: float = 2.55
    density: float = 820.0
    heat_capacity: float = 2600.0
    thermal_conductivity: float = 0.36
    heat_transfer_coefficient: float = 25.0
    mass_transfer_coefficient: float = 8.0e-7


@dataclass(frozen=True)
class FVTopology:
    """顶点中心有限体积网格的体积、内表面与对流边界信息。"""

    shape: tuple[int, ...]
    volume: np.ndarray
    face_left: np.ndarray
    face_right: np.ndarray
    face_geometry: np.ndarray
    convective_area: np.ndarray

    @property
    def size(self) -> int:
        return int(self.volume.size)


def diffusivity(moisture: np.ndarray) -> np.ndarray:
    """按附件2计算水分扩散系数，单位为 m²/s。"""
    moisture = np.asarray(moisture, dtype=float)
    if np.any(moisture <= 0.0):
        raise ArithmeticError("计算得到非正含水率，无法继续计算扩散系数。")
    return 7.0e-9 * np.exp(-0.89 / moisture)


def load_boundary(path: Path, end_time: float) -> tuple[np.ndarray, ...]:
    """读取附件1并截取从零时刻到复核终点的边界数据。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(
            min_row=2, min_col=1, max_col=3, values_only=True
        ):
            if all(value is None for value in row):
                continue
            if any(value is None for value in row):
                raise ValueError("附件1存在不完整的数据行。")
            rows.append(tuple(float(value) for value in row))
    finally:
        workbook.close()

    data = np.asarray(rows, dtype=float)
    if data.size == 0 or not np.all(np.isfinite(data)):
        raise ValueError("附件1中没有有效的有限数值数据。")
    data = data[np.argsort(data[:, 0])]
    data = data[(data[:, 0] >= 0.0) & (data[:, 0] <= end_time)]
    if data.shape[0] < 2 or np.any(np.diff(data[:, 0]) <= 0.0):
        raise ValueError("附件1的时间必须唯一且严格递增。")
    if not np.isclose(data[0, 0], 0.0) or not np.isclose(
        data[-1, 0], end_time
    ):
        raise ValueError(f"附件1必须完整覆盖 0--{end_time:g} s。")
    return data[:, 0], data[:, 1], data[:, 2]


def cell_bounds(
    nodes: np.ndarray, lower: float, upper: float
) -> tuple[np.ndarray, np.ndarray]:
    """返回顶点中心控制体在指定方向上的两侧边界。"""
    spacing = float(nodes[1] - nodes[0])
    low = np.maximum(lower, nodes - 0.5 * spacing)
    high = np.minimum(upper, nodes + 0.5 * spacing)
    return low, high


def make_1d_topology(
    params: Parameters, dr: float
) -> tuple[np.ndarray, FVTopology]:
    """建立只考虑侧面对流的一维径向有限体积网格。"""
    intervals = round(params.radius / dr)
    if intervals < 1 or not np.isclose(intervals * dr, params.radius):
        raise ValueError("--dr 必须能够整除 R=0.02 m。")

    radius = np.linspace(0.0, params.radius, intervals + 1)
    r_low, r_high = cell_bounds(radius, 0.0, params.radius)
    volume = np.pi * params.length * (r_high**2 - r_low**2)
    face_radius = 0.5 * (radius[:-1] + radius[1:])
    face_geometry = 2.0 * np.pi * params.length * face_radius / dr
    face_left = np.arange(radius.size - 1, dtype=int)
    face_right = face_left + 1
    convective_area = np.zeros(radius.size, dtype=float)
    convective_area[-1] = 2.0 * np.pi * params.radius * params.length
    topology = FVTopology(
        shape=(radius.size,),
        volume=volume,
        face_left=face_left,
        face_right=face_right,
        face_geometry=face_geometry,
        convective_area=convective_area,
    )
    return radius, topology


def make_2d_topology(
    params: Parameters, dr: float, dz: float
) -> tuple[np.ndarray, np.ndarray, FVTopology]:
    """建立 z=0 为端面、z=H 为对称中面的半圆柱网格。"""
    radial_intervals = round(params.radius / dr)
    half_length = 0.5 * params.length
    axial_intervals = round(half_length / dz)
    if radial_intervals < 1 or not np.isclose(
        radial_intervals * dr, params.radius
    ):
        raise ValueError("--dr 必须能够整除 R=0.02 m。")
    if axial_intervals < 1 or not np.isclose(
        axial_intervals * dz, half_length
    ):
        raise ValueError("--dz 必须能够整除 H=0.125 m。")

    radius = np.linspace(0.0, params.radius, radial_intervals + 1)
    axial_position = np.linspace(0.0, half_length, axial_intervals + 1)
    r_low, r_high = cell_bounds(radius, 0.0, params.radius)
    z_low, z_high = cell_bounds(axial_position, 0.0, half_length)
    radial_section = np.pi * (r_high**2 - r_low**2)
    axial_width = z_high - z_low
    volume = radial_section[:, None] * axial_width[None, :]

    nr, nz = radius.size, axial_position.size
    node_index = np.arange(nr * nz, dtype=int).reshape(nr, nz)

    radial_face_radius = 0.5 * (radius[:-1] + radius[1:])
    radial_left = node_index[:-1, :].ravel()
    radial_right = node_index[1:, :].ravel()
    radial_geometry = (
        2.0
        * np.pi
        * radial_face_radius[:, None]
        * axial_width[None, :]
        / dr
    ).ravel()

    axial_left = node_index[:, :-1].ravel()
    axial_right = node_index[:, 1:].ravel()
    axial_geometry = np.repeat(
        radial_section[:, None] / dz, nz - 1, axis=1
    ).ravel()

    convective_area = np.zeros((nr, nz), dtype=float)
    # 圆柱侧面 r=R 施加对流边界。
    convective_area[-1, :] += 2.0 * np.pi * params.radius * axial_width
    # 物理端面 z=0 施加对流边界；中面 z=H 保持对称。
    convective_area[:, 0] += radial_section

    topology = FVTopology(
        shape=(nr, nz),
        volume=volume.ravel(),
        face_left=np.concatenate((radial_left, axial_left)),
        face_right=np.concatenate((radial_right, axial_right)),
        face_geometry=np.concatenate((radial_geometry, axial_geometry)),
        convective_area=convective_area.ravel(),
    )
    return radius, axial_position, topology


def assemble_matrix(
    storage: np.ndarray,
    boundary_conductance: np.ndarray,
    topology: FVTopology,
    face_conductance: np.ndarray,
) -> csc_matrix:
    """组装储存项、边界项及守恒型内表面通量矩阵。"""
    diagonal = storage + boundary_conductance
    np.add.at(diagonal, topology.face_left, face_conductance)
    np.add.at(diagonal, topology.face_right, face_conductance)
    node_index = np.arange(topology.size, dtype=int)
    rows = np.concatenate((node_index, topology.face_left, topology.face_right))
    columns = np.concatenate((node_index, topology.face_right, topology.face_left))
    values = np.concatenate((diagonal, -face_conductance, -face_conductance))
    return coo_matrix(
        (values, (rows, columns)), shape=(topology.size, topology.size)
    ).tocsc()


def harmonic_face_value(field: np.ndarray, topology: FVTopology) -> np.ndarray:
    """由相邻节点值计算内表面的调和平均。"""
    left = field[topology.face_left]
    right = field[topology.face_right]
    return 2.0 * left * right / (left + right)


def make_heat_solver(
    params: Parameters, topology: FVTopology, dt: float
) -> tuple[Callable[[np.ndarray], np.ndarray], np.ndarray, np.ndarray]:
    """预分解后向 Euler 温度方程的常系数矩阵。"""
    storage = params.density * params.heat_capacity * topology.volume / dt
    boundary = params.heat_transfer_coefficient * topology.convective_area
    faces = params.thermal_conductivity * topology.face_geometry
    solver = factorized(assemble_matrix(storage, boundary, topology, faces))
    return solver, storage, boundary


def solve_heat_step(
    old: np.ndarray,
    ambient: float,
    solver: Callable[[np.ndarray], np.ndarray],
    storage: np.ndarray,
    boundary: np.ndarray,
) -> np.ndarray:
    """推进一个后向 Euler 温度时间层。"""
    return np.asarray(solver(storage * old + boundary * ambient), dtype=float)


def solve_moisture_step(
    old: np.ndarray,
    ambient: float,
    params: Parameters,
    topology: FVTopology,
    dt: float,
    tolerance: float,
    max_iterations: int,
) -> tuple[np.ndarray, int]:
    """用 Picard 迭代推进一个非线性水分时间层。"""
    storage = topology.volume / dt
    boundary = params.mass_transfer_coefficient * topology.convective_area
    guess = old.copy()
    for iteration in range(1, max_iterations + 1):
        node_diffusivity = diffusivity(guess)
        face_diffusivity = harmonic_face_value(node_diffusivity, topology)
        matrix = assemble_matrix(
            storage,
            boundary,
            topology,
            face_diffusivity * topology.face_geometry,
        )
        updated = np.asarray(
            spsolve(matrix, storage * old + boundary * ambient), dtype=float
        )
        error = float(np.max(np.abs(updated - guess)))
        guess = updated
        if error <= tolerance:
            return updated, iteration
    raise RuntimeError(f"水分 Picard 迭代在 {max_iterations} 次内未收敛。")


def axial_average(field: np.ndarray, axial_position: np.ndarray) -> np.ndarray:
    """按轴向控制体宽度计算二维场的轴向平均值。"""
    z_low, z_high = cell_bounds(
        axial_position, float(axial_position[0]), float(axial_position[-1])
    )
    return np.average(field, axis=1, weights=z_high - z_low)


def influence_length(
    field: np.ndarray,
    axial_position: np.ndarray,
    fraction: float,
    absolute_tolerance: float,
) -> tuple[float, np.ndarray]:
    """计算端面扰动衰减至给定比例后的单侧影响长度。"""
    deviation = np.max(np.abs(field - field[:, [-1]]), axis=0)
    if deviation[0] <= absolute_tolerance:
        return 0.0, deviation
    threshold = max(fraction * deviation[0], absolute_tolerance)
    tail_maximum = np.maximum.accumulate(deviation[::-1])[::-1]
    candidates = np.flatnonzero(tail_maximum <= threshold)
    length = (
        float(axial_position[candidates[0]])
        if candidates.size
        else float(axial_position[-1])
    )
    return length, deviation


def positive_for_log(values: np.ndarray) -> np.ndarray:
    """将精确零值隐藏，使半对数坐标能够保留初始时刻。"""
    values = np.asarray(values, dtype=float)
    return np.where(values > 0.0, values, np.nan)


def parse_arguments(script_dir: Path) -> argparse.Namespace:
    """读取二维复核与图8输出参数。"""
    parser = argparse.ArgumentParser(
        description="计算问题1的一维/二维对照并绘制二维轴对称复核图。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=script_dir.parent / "附件" / "附件1.xlsx",
        help="附件1路径",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=script_dir / "图",
        help="图8输出目录",
    )
    parser.add_argument("--end-time", type=float, default=1800.0)
    parser.add_argument("--dt", type=float, default=0.25)
    parser.add_argument("--dr", type=float, default=0.00025)
    parser.add_argument("--dz", type=float, default=0.0005)
    parser.add_argument("--save-interval", type=float, default=60.0)
    parser.add_argument("--picard-tolerance", type=float, default=1.0e-11)
    parser.add_argument("--max-picard-iterations", type=int, default=50)
    parser.add_argument(
        "--influence-fraction",
        type=float,
        default=0.05,
        help="定义端部影响长度时采用的剩余轴向偏离比例",
    )
    return parser.parse_args()


def validate_arguments(args: argparse.Namespace) -> tuple[int, int]:
    """检查时间、网格和端部影响判据是否合法。"""
    if min(args.end_time, args.dt, args.dr, args.dz, args.save_interval) <= 0.0:
        raise ValueError("时间和网格参数必须为正数。")
    if not 0.0 < args.influence_fraction < 1.0:
        raise ValueError("--influence-fraction 必须位于 0 与 1 之间。")

    step_count = round(args.end_time / args.dt)
    save_stride = round(args.save_interval / args.dt)
    if not np.isclose(step_count * args.dt, args.end_time):
        raise ValueError("--dt 必须能够整除 --end-time。")
    if save_stride < 1 or not np.isclose(
        save_stride * args.dt, args.save_interval
    ):
        raise ValueError("--dt 必须能够整除 --save-interval。")
    return step_count, save_stride


def run_validation(args: argparse.Namespace) -> dict[str, np.ndarray | float]:
    """推进一维与二维模型，并返回绘图所需的完整精度状态。"""
    step_count, save_stride = validate_arguments(args)
    params = Parameters()
    source_time, source_temperature, source_moisture = load_boundary(
        args.input.resolve(), args.end_time
    )
    radius, topology_1d = make_1d_topology(params, args.dr)
    radius_2d, axial_position, topology_2d = make_2d_topology(
        params, args.dr, args.dz
    )
    if not np.array_equal(radius, radius_2d):
        raise ArithmeticError("一维与二维模型的径向网格不一致。")

    solve_t1, storage_t1, boundary_t1 = make_heat_solver(
        params, topology_1d, args.dt
    )
    solve_t2, storage_t2, boundary_t2 = make_heat_solver(
        params, topology_2d, args.dt
    )

    temperature_1d = np.full(topology_1d.size, params.initial_temperature)
    moisture_1d = np.full(topology_1d.size, params.initial_moisture)
    temperature_2d = np.full(topology_2d.size, params.initial_temperature)
    moisture_2d = np.full(topology_2d.size, params.initial_moisture)

    saved_time: list[float] = []
    saved_temperature_1d: list[np.ndarray] = []
    saved_moisture_1d: list[np.ndarray] = []
    saved_temperature_2d: list[np.ndarray] = []
    saved_moisture_2d: list[np.ndarray] = []
    maximum_picard_1d = 0
    maximum_picard_2d = 0

    temperature_span = max(
        1.0e-14,
        float(np.max(np.abs(source_temperature - params.initial_temperature))),
    )
    moisture_span = max(
        1.0e-14,
        float(params.initial_moisture - np.min(source_moisture)),
    )

    def save_state(time_s: float) -> None:
        saved_time.append(time_s)
        saved_temperature_1d.append(temperature_1d.copy())
        saved_moisture_1d.append(moisture_1d.copy())
        saved_temperature_2d.append(
            temperature_2d.reshape(topology_2d.shape).copy()
        )
        saved_moisture_2d.append(moisture_2d.reshape(topology_2d.shape).copy())

    save_state(0.0)
    for step in range(1, step_count + 1):
        time_s = step * args.dt
        ambient_temperature = float(
            np.interp(time_s, source_time, source_temperature)
        )
        ambient_moisture = float(np.interp(time_s, source_time, source_moisture))

        temperature_1d = solve_heat_step(
            temperature_1d,
            ambient_temperature,
            solve_t1,
            storage_t1,
            boundary_t1,
        )
        temperature_2d = solve_heat_step(
            temperature_2d,
            ambient_temperature,
            solve_t2,
            storage_t2,
            boundary_t2,
        )
        moisture_1d, iterations_1d = solve_moisture_step(
            moisture_1d,
            ambient_moisture,
            params,
            topology_1d,
            args.dt,
            args.picard_tolerance,
            args.max_picard_iterations,
        )
        moisture_2d, iterations_2d = solve_moisture_step(
            moisture_2d,
            ambient_moisture,
            params,
            topology_2d,
            args.dt,
            args.picard_tolerance,
            args.max_picard_iterations,
        )
        maximum_picard_1d = max(maximum_picard_1d, iterations_1d)
        maximum_picard_2d = max(maximum_picard_2d, iterations_2d)

        if step % save_stride == 0:
            save_state(time_s)
            print(f"已保存 t={time_s:g} s / {args.end_time:g} s")

    temperature_2d_history = np.asarray(saved_temperature_2d)
    moisture_2d_history = np.asarray(saved_moisture_2d)
    if not np.all(np.isfinite(temperature_2d_history)) or not np.all(
        np.isfinite(moisture_2d_history)
    ):
        raise ArithmeticError("二维计算场中出现 NaN 或无穷值。")

    print(
        "最大 Picard 迭代次数："
        f"一维={maximum_picard_1d}，二维={maximum_picard_2d}"
    )
    return {
        "time_s": np.asarray(saved_time, dtype=float),
        "radius_m": radius,
        "axial_position_m": axial_position,
        "temperature_1d": np.asarray(saved_temperature_1d),
        "moisture_1d": np.asarray(saved_moisture_1d),
        "temperature_2d": temperature_2d_history,
        "moisture_2d": moisture_2d_history,
        "influence_fraction": float(args.influence_fraction),
        "temperature_span": temperature_span,
        "moisture_span": moisture_span,
    }


def draw_validation_figure(
    data: dict[str, np.ndarray | float], output_dir: Path
) -> Path:
    """根据完整精度状态绘制端面衰减和误差随时间变化。"""
    time_s = np.asarray(data["time_s"])
    axial_position_cm = 100.0 * np.asarray(data["axial_position_m"])
    temperature_1d = np.asarray(data["temperature_1d"])
    moisture_1d = np.asarray(data["moisture_1d"])
    temperature_2d = np.asarray(data["temperature_2d"])
    moisture_2d = np.asarray(data["moisture_2d"])
    fraction = float(data["influence_fraction"])
    temperature_span = float(data["temperature_span"])
    moisture_span = float(data["moisture_span"])

    temperature_difference = np.abs(
        temperature_2d - temperature_1d[:, :, None]
    )
    moisture_difference = np.abs(moisture_2d - moisture_1d[:, :, None])
    worst_temperature_index = int(
        np.argmax(np.max(temperature_difference, axis=(1, 2)))
    )
    worst_moisture_index = int(
        np.argmax(np.max(moisture_difference, axis=(1, 2)))
    )

    temperature_length, axial_temperature = influence_length(
        temperature_2d[worst_temperature_index],
        axial_position_cm,
        fraction,
        1.0e-10,
    )
    moisture_length, axial_moisture = influence_length(
        moisture_2d[worst_moisture_index],
        axial_position_cm,
        fraction,
        1.0e-12,
    )
    normalized_temperature = (
        axial_temperature / axial_temperature[0]
        if axial_temperature[0] > 0.0
        else np.zeros_like(axial_temperature)
    )
    normalized_moisture = (
        axial_moisture / axial_moisture[0]
        if axial_moisture[0] > 0.0
        else np.zeros_like(axial_moisture)
    )

    middle_temperature_error = np.max(
        np.abs(temperature_2d[:, :, -1] - temperature_1d), axis=1
    )
    middle_moisture_error = np.max(
        np.abs(moisture_2d[:, :, -1] - moisture_1d), axis=1
    )
    average_temperature = np.asarray(
        [axial_average(field, axial_position_cm) for field in temperature_2d]
    )
    average_moisture = np.asarray(
        [axial_average(field, axial_position_cm) for field in moisture_2d]
    )
    average_temperature_error = np.max(
        np.abs(average_temperature - temperature_1d), axis=1
    )
    average_moisture_error = np.max(
        np.abs(average_moisture - moisture_1d), axis=1
    )

    plt.rcParams.update(
        {
            "font.sans-serif": [
                "Microsoft YaHei",
                "SimHei",
                "Noto Sans CJK SC",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
            "font.size": 14,
        }
    )
    fig, (axis_decay, axis_time) = plt.subplots(
        1, 2, figsize=(16.0, 6.4), constrained_layout=True
    )

    axis_decay.plot(
        axial_position_cm,
        normalized_moisture,
        linewidth=2.6,
        label="含水率轴向偏离",
    )
    axis_decay.plot(
        axial_position_cm,
        normalized_temperature,
        linewidth=2.6,
        label="温度轴向偏离",
    )
    axis_decay.axhline(
        fraction,
        color="#d62728",
        linestyle="--",
        linewidth=1.8,
        label=f"剩余偏离 {100 * fraction:g}%",
    )
    axis_decay.axvline(moisture_length, color="C0", linestyle=":", linewidth=1.4)
    axis_decay.axvline(
        temperature_length, color="C1", linestyle=":", linewidth=1.4
    )
    axis_decay.set_xlim(
        float(axial_position_cm[0]), float(axial_position_cm[-1])
    )
    axis_decay.set_ylim(-0.02, 1.03)
    axis_decay.set_xlabel("距端面的轴向距离 z / cm")
    axis_decay.set_ylabel("归一化轴向偏离")
    axis_decay.set_title("(a) 端面影响沿轴向的衰减")
    axis_decay.grid(alpha=0.25)
    axis_decay.legend(loc="upper right", frameon=True, fontsize=13)
    axis_decay.text(
        0.98,
        0.54,
        f"含水率影响长度：{moisture_length:.2f} cm\n"
        f"温度影响长度：{temperature_length:.2f} cm",
        transform=axis_decay.transAxes,
        ha="right",
        va="top",
        bbox=dict(facecolor="white", edgecolor="0.7", alpha=0.9),
    )

    time_minute = time_s / 60.0
    middle_moisture_percent = 100.0 * middle_moisture_error / moisture_span
    average_moisture_percent = 100.0 * average_moisture_error / moisture_span
    middle_temperature_percent = (
        100.0 * middle_temperature_error / temperature_span
    )
    average_temperature_percent = (
        100.0 * average_temperature_error / temperature_span
    )
    middle_display_factor = 1.0e4

    axis_time.semilogy(
        time_minute,
        positive_for_log(middle_moisture_percent * middle_display_factor),
        linewidth=2.6,
        label=f"含水率中面差异（显示×{middle_display_factor:.0e}）",
    )
    axis_time.semilogy(
        time_minute,
        positive_for_log(average_moisture_percent),
        linewidth=2.6,
        linestyle="--",
        label="含水率轴向平均差异",
    )
    axis_time.semilogy(
        time_minute,
        positive_for_log(middle_temperature_percent * middle_display_factor),
        linewidth=2.6,
        label=f"温度中面差异（显示×{middle_display_factor:.0e}）",
    )
    axis_time.semilogy(
        time_minute,
        positive_for_log(average_temperature_percent),
        linewidth=2.6,
        linestyle="--",
        label="温度轴向平均差异",
    )
    axis_time.set_xlabel("时间 / min")
    axis_time.set_ylabel("相对于完整变化跨度的差异 / %")
    axis_time.set_title("(b) 中面一致性与轴向平均差异随时间变化")
    axis_time.set_ylim(1.0e-12, 5.0)
    axis_time.grid(alpha=0.25)
    axis_time.legend(loc="lower right", frameon=True, fontsize=12)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "图8 一维径向模型的二维轴对称复核.pdf"
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)
    print(f"含水率端面影响长度：{moisture_length:.4f} cm")
    print(f"温度端面影响长度：{temperature_length:.4f} cm")
    print(f"图8已保存至：{output_path.resolve()}")
    return output_path


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    arguments = parse_arguments(script_dir)
    validation_data = run_validation(arguments)
    draw_validation_figure(validation_data, arguments.output_dir)


if __name__ == "__main__":
    main()
